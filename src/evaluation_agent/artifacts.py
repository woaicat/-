"""Validate and atomically register files produced by one authorized turn."""

import csv
import json
import shutil
import tempfile
from io import StringIO
from pathlib import Path, PurePosixPath
from uuid import uuid4

from docx import Document

from evaluation_agent.store import ConflictError, NotFoundError, WorkspaceStore, _now


MAX_SOURCE_BYTES = 10 * 1024 * 1024


def _validate_draft_path(turn_id: str, draft_path: str) -> None:
    if not draft_path.startswith("/") or ".." in draft_path.split("/"):
        raise ValueError("Invalid draft path")
    if not str(PurePosixPath(draft_path)).startswith(f"/drafts/{turn_id}/"):
        raise ValueError("Draft belongs to another turn")


def _validate_filename(filename: str) -> str:
    if not filename or filename in (".", "..") or "/" in filename or "\\" in filename:
        raise ValueError("Invalid artifact filename")
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix not in ("md", "csv", "docx"):
        raise ValueError("Unsupported artifact type")
    return suffix


def _write_docx(markdown: str, destination: Path) -> None:
    """Keep a Markdown source and export readable headings, bullets and tables."""
    document = Document()
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            if level <= 4 and line[level:level + 1] == " ":
                document.add_heading(line[level:].strip(), level=min(level, 4))
                index += 1
                continue
        if line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
            index += 1
            continue
        if line.startswith("| ") and " |" in line:
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(set(cell) <= {"-", ":", " "} for cell in cells):
                    table_rows.append(cells)
                index += 1
            if table_rows:
                width = max(map(len, table_rows))
                table = document.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for cells in table_rows:
                    row = table.add_row()
                    for column, cell in enumerate(cells):
                        row.cells[column].text = cell
            continue
        document.add_paragraph(line)
        index += 1
    document.save(destination)
    if not Document(destination).paragraphs and not Document(destination).tables:
        raise ValueError("DOCX export is empty")


def _validate_csv(content: str) -> None:
    rows = list(csv.reader(StringIO(content)))
    if len(rows) < 2 or not rows[0] or not all(cell.strip() for cell in rows[0]):
        raise ValueError("CSV needs a header and at least one data row")
    width = len(rows[0])
    if any(len(row) != width for row in rows[1:]):
        raise ValueError("CSV rows have inconsistent column counts")


class ArtifactPublisher:
    def __init__(self, store: WorkspaceStore, output_root: str | Path):
        self.store = store
        self.output_root = Path(output_root).resolve()
        self.output_root.mkdir(parents=True, exist_ok=True)

    def publish(
        self,
        conversation_id: str,
        turn_id: str,
        *,
        draft_path: str,
        filename: str,
        content: str,
        deliverable_id: str | None = None,
        artifact_id: str | None = None,
    ) -> dict:
        _validate_draft_path(turn_id, draft_path)
        kind = _validate_filename(filename)
        if not content.strip() or len(content.encode("utf-8")) > MAX_SOURCE_BYTES:
            raise ValueError("Draft is empty or too large")
        if kind == "csv":
            _validate_csv(content)

        staged = Path(tempfile.mkdtemp(prefix=".staged-", dir=self.output_root))
        final_dir: Path | None = None
        published = False
        try:
            source_name = "source.csv" if kind == "csv" else "source.md"
            (staged / source_name).write_text(content, encoding="utf-8")
            export_name = filename if kind == "docx" else source_name
            if kind == "docx":
                _write_docx(content, staged / export_name)
            if not (staged / export_name).is_file():
                raise ValueError("Export was not created")
            with self.store._transaction() as conn:
                turn = conn.execute(
                    """SELECT evaluation_id,scope_kind,scope_json,outline_id,outline_revision,status
                       FROM turns WHERE conversation_id=? AND turn_id=?""",
                    (conversation_id, turn_id),
                ).fetchone()
                if turn is None:
                    raise NotFoundError("Turn not found")
                if turn["status"] != "running":
                    raise ConflictError("Turn cannot publish after cancellation or completion", turn_id=turn_id)
                if not turn["evaluation_id"]:
                    raise ConflictError("Turn has no product scope", turn_id=turn_id)
                scope = json.loads(turn["scope_json"])
                if scope.get("finalizing"):
                    raise ConflictError("Turn is finalizing and cannot publish more files", turn_id=turn_id)
                scope_kind = turn["scope_kind"]
                if scope_kind in ("package_generate", "package_continue"):
                    if not deliverable_id or deliverable_id not in scope.get("deliverable_ids", []):
                        raise ConflictError("Deliverable is outside confirmed scope", turn_id=turn_id)
                    deliverable = conn.execute(
                        """SELECT name,kind,status,artifact_id FROM deliverables WHERE deliverable_id=?
                           AND conversation_id=? AND evaluation_id=? AND outline_id=? AND outline_revision=?""",
                        (deliverable_id, conversation_id, turn["evaluation_id"],
                         turn["outline_id"], turn["outline_revision"]),
                    ).fetchone()
                    matches_name = deliverable is not None and self.store._artifact_name_matches(
                        conn, conversation_id, deliverable["artifact_id"] or "", deliverable["name"], filename,
                    )
                    if deliverable is None or not matches_name or deliverable["kind"] != kind:
                        raise ConflictError("Deliverable identity does not match the confirmation", turn_id=turn_id)
                    if deliverable["status"] == "ready":
                        if not artifact_id or artifact_id != deliverable["artifact_id"]:
                            raise ConflictError("Ready deliverable needs its current artifact ID", turn_id=turn_id)
                        first_publication = conn.execute(
                            """SELECT 1 FROM events WHERE conversation_id=? AND turn_id=?
                               AND type='artifact.ready' AND json_extract(payload_json,'$.artifact_id')=?""",
                            (conversation_id, turn_id, artifact_id),
                        ).fetchone()
                        if first_publication is None and scope_kind != "package_continue":
                            raise ConflictError("Only this turn's published file can be repaired", turn_id=turn_id)
                    elif artifact_id:
                        raise ConflictError("Pending deliverable cannot overwrite an artifact", turn_id=turn_id)
                    outline = conn.execute(
                        "SELECT status,revision FROM outlines WHERE outline_id=? AND conversation_id=?",
                        (turn["outline_id"], conversation_id),
                    ).fetchone()
                    if outline is None or outline["status"] != "confirmed" or outline["revision"] != turn["outline_revision"]:
                        raise ConflictError("Confirmed outline revision is no longer valid", turn_id=turn_id)
                elif scope_kind == "focused_create":
                    if artifact_id or deliverable_id or Path(filename).stem not in scope.get("target_names", []):
                        raise ConflictError("New file is outside focused scope", turn_id=turn_id)
                elif scope_kind == "focused_edit":
                    if deliverable_id or not artifact_id or artifact_id not in scope.get("target_artifact_ids", []):
                        raise ConflictError("Edit is outside focused scope", turn_id=turn_id)
                else:
                    raise ConflictError("Turn has no publication permission", turn_id=turn_id)

                prior = None
                if artifact_id:
                    prior = conn.execute(
                        """SELECT filename,kind,content_revision FROM artifacts WHERE artifact_id=?
                           AND conversation_id=? AND evaluation_id=?""",
                        (artifact_id, conversation_id, turn["evaluation_id"]),
                    ).fetchone()
                    if prior is None or prior["filename"] != filename or prior["kind"] != kind:
                        raise NotFoundError("Target artifact is not in this product")
                else:
                    alias = conn.execute(
                        "SELECT 1 FROM artifact_aliases WHERE conversation_id=? AND evaluation_id=? AND filename=?",
                        (conversation_id, turn["evaluation_id"], filename),
                    ).fetchone()
                    if alias:
                        raise ConflictError("Filename is reserved by an existing artifact alias", turn_id=turn_id)
                    duplicate = conn.execute(
                        "SELECT artifact_id FROM artifacts WHERE conversation_id=? AND evaluation_id=? AND filename=?",
                        (conversation_id, turn["evaluation_id"], filename),
                    ).fetchone()
                    if duplicate:
                        raise ConflictError("Artifact already exists; edit its stable ID", turn_id=turn_id)
                    artifact_id = uuid4().hex
                revision = prior["content_revision"] + 1 if prior else 1
                final_dir = self.output_root / conversation_id / turn["evaluation_id"] / artifact_id / f"rev-{revision}"
                final_dir.parent.mkdir(parents=True, exist_ok=True)
                staged.rename(final_dir)
                source_path = str(final_dir / source_name)
                export_path = str(final_dir / export_name)
                now = _now()
                if prior:
                    conn.execute(
                        """UPDATE artifacts SET source_path=?,export_path=?,content_revision=?,ready_at=?
                           WHERE artifact_id=?""",
                        (source_path, export_path, revision, now, artifact_id),
                    )
                    event_type = "artifact.updated"
                else:
                    conn.execute(
                        """INSERT INTO artifacts(artifact_id,conversation_id,evaluation_id,filename,kind,
                           source_path,export_path,content_revision,ready_at) VALUES(?,?,?,?,?,?,?,?,?)""",
                        (artifact_id, conversation_id, turn["evaluation_id"], filename, kind,
                         source_path, export_path, revision, now),
                    )
                    event_type = "artifact.ready"
                if deliverable_id:
                    conn.execute(
                        "UPDATE deliverables SET status='ready',artifact_id=? WHERE deliverable_id=?",
                        (artifact_id, deliverable_id),
                    )
                event_id = self.store._event(conn, conversation_id, turn_id, event_type, {
                    "artifact_id": artifact_id, "evaluation_id": turn["evaluation_id"],
                    "filename": filename, "kind": kind, "content_revision": revision,
                    "deliverable_id": deliverable_id,
                })
            published = True
            return {"artifact_id": artifact_id, "content_revision": revision, "event_id": event_id}
        finally:
            if staged.exists():
                shutil.rmtree(staged)
            if not published and final_dir is not None and final_dir.exists():
                shutil.rmtree(final_dir)

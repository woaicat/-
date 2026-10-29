"""Local single-user HTTP and replayable SSE contract."""

import asyncio
import json
import shutil
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from langchain_core.language_models.chat_models import BaseChatModel
from langgraph.checkpoint.sqlite.aio import AsyncSqliteSaver
from langgraph.config import get_config
from pydantic import BaseModel, Field

from evaluation_agent.artifacts import ArtifactPublisher
from evaluation_agent.executor import TurnExecutor
from evaluation_agent.observability import configure_observability, flush_traces
from evaluation_agent.product_tools import create_product_tools
from evaluation_agent.runtime import ContextBudget, create_runtime
from evaluation_agent.scoped_backend import ScopedStateBackend
from evaluation_agent.store import ConflictError, NotFoundError, WorkspaceStore


MAX_UPLOAD_BYTES = 10 * 1024 * 1024


class ConversationInput(BaseModel):
    title: str = "新对话"


class MessageInput(BaseModel):
    body: str = ""
    client_key: str
    attachment_ids: list[str] = Field(default_factory=list)


class OutlineEditInput(BaseModel):
    expected_revision: int
    body: str


class ConfirmInput(BaseModel):
    expected_revision: int
    client_key: str


def _safe_snapshot(snapshot: dict) -> dict:
    for attachment in snapshot["attachments"]:
        attachment.pop("storage_path", None)
        attachment.pop("extracted_path", None)
    for artifact in snapshot["artifacts"]:
        artifact.pop("source_path", None)
        artifact.pop("export_path", None)
        artifact["preview_url"] = (
            f"/api/conversations/{artifact['conversation_id']}/artifacts/{artifact['artifact_id']}/preview"
        )
    return snapshot


def _extract_docx(raw: bytes) -> str:
    try:
        document = Document(BytesIO(raw))
    except Exception as exc:
        raise ValueError("无法解析 Word 文档") from exc
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    if not lines:
        raise ValueError("Word 文档没有可读取的文字")
    return "\n".join(lines)


def create_app(
    *,
    data_root: str | Path | None = None,
    model: BaseChatModel | None = None,
    model_provider: str | None = None,
    budget: ContextBudget = ContextBudget(),
) -> FastAPI:
    root = Path(data_root or Path(__file__).resolve().parents[2] / "storage").resolve()
    root.mkdir(parents=True, exist_ok=True)
    store = WorkspaceStore(root / "workspace.sqlite")
    publisher = ArtifactPublisher(store, root / "outputs")

    @asynccontextmanager
    async def lifespan(app: FastAPI):
        configure_observability()
        store.recover_interrupted_turns()
        app.state.executor = None
        if model is None:
            yield
            return
        if not model_provider:
            raise RuntimeError("model_provider is required with a model")
        async with AsyncSqliteSaver.from_conn_string(str(root / "checkpoints.sqlite")) as saver:
            backend = ScopedStateBackend()
            tools = create_product_tools(store, publisher, backend)
            def observe_context(kind: str, payload: dict) -> None:
                config = get_config().get("configurable", {})
                store.record_context_event(config["conversation_id"], config["turn_id"], kind, payload)

            agent = create_runtime(
                model=model, model_provider=model_provider, product_tools=tools,
                checkpointer=saver, backend=backend, budget=budget, context_observer=observe_context,
            )
            executor = TurnExecutor(store, agent)
            app.state.executor = executor
            try:
                yield
            finally:
                await executor.shutdown()
                await flush_traces()

    app = FastAPI(title="对话式评测方案智能体", lifespan=lifespan)
    app.state.store = store
    frontend = Path(__file__).resolve().parents[2] / "frontend"
    app.mount("/assets", StaticFiles(directory=frontend), name="assets")

    @app.get("/", include_in_schema=False)
    async def frontend_home():
        assets = (
            frontend / "index.html", frontend / "app.js", frontend / "styles.css",
            frontend / "vendor" / "markdown-it.umd.min.js",
            frontend / "vendor" / "purify.min.js",
        )
        version = max(path.stat().st_mtime_ns for path in assets)
        html = (frontend / "index.html").read_text(encoding="utf-8")
        return HTMLResponse(html.replace("{{ASSET_VERSION}}", str(version)), headers={
            "Cache-Control": "no-store",
        })

    @app.exception_handler(NotFoundError)
    async def not_found(_, exc: NotFoundError):
        return JSONResponse(status_code=404, content={"detail": str(exc)})

    @app.exception_handler(ConflictError)
    async def conflict(_, exc: ConflictError):
        return JSONResponse(status_code=409, content={"detail": str(exc), "turn_id": exc.turn_id})

    @app.exception_handler(ValueError)
    async def invalid(_, exc: ValueError):
        return JSONResponse(status_code=422, content={"detail": str(exc)})

    def require_executor(request: Request) -> TurnExecutor:
        executor = request.app.state.executor
        if executor is None:
            raise HTTPException(status_code=503, detail="模型服务尚未配置")
        return executor

    @app.post("/api/conversations")
    async def create_conversation(payload: ConversationInput):
        return store.create_conversation(payload.title)

    @app.get("/api/conversations")
    async def list_conversations():
        return store.list_conversations()

    @app.get("/api/conversations/{conversation_id}")
    async def get_conversation(conversation_id: str):
        snapshot = store.snapshot(conversation_id)
        snapshot["context"].update({
            "window_tokens": budget.window_tokens,
            "summarize_at_tokens": budget.summarize_at_tokens,
            "max_output_tokens": budget.max_output_tokens,
        })
        completed = {
            event["payload"]["message_id"]: event["payload"].get("artifact_ids", [])
            for event in store.list_events(conversation_id)
            if event["type"] == "message.completed"
        }
        for message in snapshot["messages"]:
            if message["role"] == "assistant":
                message["artifact_ids"] = completed.get(message["message_id"], [])
        return _safe_snapshot(snapshot)

    @app.post("/api/conversations/{conversation_id}/attachments")
    async def upload_attachment(
        conversation_id: str,
        file: UploadFile = File(...),
        upload_key: str = Form(...),
    ):
        store.snapshot(conversation_id)
        if not upload_key.strip():
            raise HTTPException(status_code=422, detail="上传幂等键不能为空")
        filename = file.filename or ""
        if not filename or filename != Path(filename).name:
            raise HTTPException(status_code=422, detail="无效的文件名")
        suffix = Path(filename).suffix.lower()
        if suffix not in (".md", ".docx"):
            raise HTTPException(status_code=422, detail="目前只支持 Markdown 和 Word（.docx）")
        raw = await file.read(MAX_UPLOAD_BYTES + 1)
        if not raw or len(raw) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=422, detail="文件为空或超过 10 MB")
        if suffix == ".md":
            try:
                extracted = raw.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise HTTPException(status_code=422, detail="Markdown 文件须为 UTF-8") from exc
        else:
            extracted = _extract_docx(raw)
        upload_dir = root / "uploads" / conversation_id / uuid4().hex
        upload_dir.mkdir(parents=True, exist_ok=True)
        original = upload_dir / f"original{suffix}"
        text = upload_dir / "extracted.md"
        original.write_bytes(raw)
        text.write_text(extracted, encoding="utf-8")
        try:
            result = store.create_attachment(
                conversation_id, upload_key=upload_key, filename=filename,
                mime_type="text/markdown" if suffix == ".md" else
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                storage_path=str(original), extracted_path=str(text),
            )
        finally:
            # Existing idempotent uploads use their original persisted files.
            if "result" not in locals() or result["storage_path"] != str(original):
                shutil.rmtree(upload_dir, ignore_errors=True)
        return {key: result[key] for key in (
            "attachment_id", "conversation_id", "evaluation_id", "message_id",
            "filename", "mime_type", "created_at",
        )}

    @app.post("/api/conversations/{conversation_id}/messages")
    async def send_message(conversation_id: str, payload: MessageInput, request: Request):
        executor = require_executor(request)
        result = store.begin_message_turn(
            conversation_id, body=payload.body, client_key=payload.client_key,
            attachment_ids=payload.attachment_ids,
        )
        turn = store.get_turn(conversation_id, result["turn_id"])
        if turn["status"] == "running":
            executor.start(conversation_id, result["turn_id"], payload.body)
        return result

    @app.put("/api/conversations/{conversation_id}/outlines/{outline_id}")
    async def edit_outline(conversation_id: str, outline_id: str, payload: OutlineEditInput):
        revision = store.edit_outline(
            conversation_id, outline_id,
            expected_revision=payload.expected_revision, body=payload.body,
        )
        return {"outline_id": outline_id, "revision": revision}

    @app.post("/api/conversations/{conversation_id}/outlines/{outline_id}/confirm")
    async def confirm_outline(conversation_id: str, outline_id: str, payload: ConfirmInput, request: Request):
        executor = require_executor(request)
        result = store.confirm_outline(
            conversation_id, outline_id,
            expected_revision=payload.expected_revision, client_key=payload.client_key,
        )
        turn = store.get_turn(conversation_id, result["turn_id"])
        if turn["status"] == "running":
            snapshot = store.snapshot(conversation_id)
            outline = next(row for row in snapshot["outlines"] if row["outline_id"] == outline_id)
            executor.start(
                conversation_id, result["turn_id"],
                "用户已在大纲卡片中确认以下方案，请按确认的交付清单完成：\n" + outline["body"],
            )
        return result

    @app.post("/api/conversations/{conversation_id}/turns/{turn_id}/cancel")
    async def cancel_turn(conversation_id: str, turn_id: str, request: Request):
        executor = require_executor(request)
        status = store.request_cancel(conversation_id, turn_id)
        if status == "cancel_requested":
            if turn_id in executor.tasks:
                executor.cancel(turn_id)
            else:
                store.finish_turn(conversation_id, turn_id, "cancelled")
                status = "cancelled"
        return {"turn_id": turn_id, "status": status}

    @app.get("/api/conversations/{conversation_id}/events")
    async def stream_events(conversation_id: str, request: Request, after: int = 0):
        store.snapshot(conversation_id)
        header = request.headers.get("last-event-id")
        if header:
            try:
                after = max(after, int(header))
            except ValueError:
                raise HTTPException(status_code=422, detail="Invalid Last-Event-ID") from None

        async def generate():
            cursor = after
            while not await request.is_disconnected():
                rows = store.list_events(conversation_id, after=cursor)
                for row in rows:
                    cursor = row["event_id"]
                    envelope = {key: row[key] for key in (
                        "v", "event_id", "conversation_id", "turn_id", "created_at", "type", "payload",
                    )}
                    yield f"id: {cursor}\nevent: {row['type']}\ndata: {json.dumps(envelope, ensure_ascii=False)}\n\n"
                if not rows:
                    await asyncio.sleep(0.15)

        return StreamingResponse(generate(), media_type="text/event-stream", headers={"Cache-Control": "no-cache"})

    @app.get("/api/conversations/{conversation_id}/evaluations/{evaluation_id}/artifacts")
    async def list_artifacts(conversation_id: str, evaluation_id: str):
        rows = store.list_artifacts(conversation_id, evaluation_id)
        return [{
            "artifact_id": row["artifact_id"], "evaluation_id": evaluation_id,
            "filename": row["filename"], "kind": row["kind"],
            "content_revision": row["content_revision"], "ready_at": row["ready_at"],
            "preview_url": f"/api/conversations/{conversation_id}/artifacts/{row['artifact_id']}/preview",
        } for row in rows]

    @app.get("/api/conversations/{conversation_id}/artifacts/{artifact_id}/preview")
    async def preview_artifact(conversation_id: str, artifact_id: str):
        artifact = store.get_artifact(conversation_id, artifact_id)
        source = Path(artifact["source_path"])
        if not source.is_file():
            raise HTTPException(status_code=500, detail="产物源文件缺失")
        return FileResponse(source, media_type="text/plain; charset=utf-8")

    @app.get("/api/conversations/{conversation_id}/artifacts/{artifact_id}/download")
    async def download_artifact(conversation_id: str, artifact_id: str):
        artifact = store.get_artifact(conversation_id, artifact_id)
        exported = Path(artifact["export_path"])
        if not exported.is_file():
            raise HTTPException(status_code=500, detail="产物文件缺失")
        return FileResponse(exported, filename=artifact["filename"])

    return app

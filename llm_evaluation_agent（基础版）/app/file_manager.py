"""Helpers for managing generated artefact files."""

from __future__ import annotations

import csv
import json
import re
import shutil
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable, List, Sequence

from docx import Document

from .config import get_settings
from .models import GeneratedFile, Manifest

SAFE_FILENAME_PATTERN = re.compile(r"[^0-9a-zA-Z\u4e00-\u9fff._-]+")


def sanitize_filename(name: str, fallback: str = "file") -> str:
    """Return a filesystem-safe filename while保留扩展名."""

    candidate = Path(name or fallback).name.replace("\x00", "")
    if not candidate:
        candidate = fallback
    stem, suffix = Path(candidate).stem, Path(candidate).suffix
    safe_stem = SAFE_FILENAME_PATTERN.sub("_", stem).strip("._ ") or fallback
    safe_suffix = SAFE_FILENAME_PATTERN.sub("", suffix) if suffix else ""
    result = (safe_stem + safe_suffix)[:120]
    return result or fallback


def sanitize_path_fragment(value: str, fallback: str = "dir") -> str:
    """Sanitize directory fragments to prevent路径穿越."""

    fragment = Path(value or fallback).name.replace("\x00", "")
    fragment = SAFE_FILENAME_PATTERN.sub("_", fragment).strip("._ ")
    fragment = fragment[:80] or fallback
    return fragment


def _directory_name(product_name: str, request_id: str) -> str:
    safe_name = sanitize_path_fragment(product_name, fallback="product")
    return f"{safe_name}_{request_id}"


def create_request_directory(request_id: str, product_name: str) -> Path:
    """Create the output directory for a specific request."""

    settings = get_settings()
    directory = settings.outputs_dir / _directory_name(product_name, request_id)
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def find_request_directory(request_id: str) -> Path | None:
    """Locate the output directory for a given request id."""

    settings = get_settings()
    suffix = f"_{request_id}"
    for item in settings.outputs_dir.iterdir():
        if item.is_dir() and item.name.endswith(suffix):
            return item
    return None


def write_markdown(directory: Path, filename: str, content: str) -> Path:
    """Persist markdown content as a Word document."""

    safe_name = sanitize_filename(filename)
    docx_name = Path(safe_name).with_suffix(".docx").name
    path = directory / docx_name
    document = Document()
    for kind, payload in _parse_markdown_blocks(content or ""):
        if kind == "heading":
            level, text = payload
            document.add_heading(text, level=level)
        elif kind == "list":
            document.add_paragraph(payload, style="List Bullet")
        elif kind == "numbered":
            document.add_paragraph(payload, style="List Number")
        elif kind == "table":
            rows = payload
            col_count = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=col_count)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx in range(col_count):
                    cell_value = row[c_idx] if c_idx < len(row) else ""
                    table.cell(r_idx, c_idx).text = cell_value
        else:
            document.add_paragraph(payload)
    document.save(path)
    return path


def write_csv(directory: Path, filename: str, rows: Sequence[Sequence[str]]) -> Path:
    """Write CSV content."""

    safe_name = sanitize_filename(filename)
    path = directory / safe_name
    with path.open("w", encoding="utf-8", newline="") as csv_file:
        writer = csv.writer(csv_file)
        for row in rows:
            writer.writerow(row)
    return path


def write_text(directory: Path, filename: str, content: str) -> Path:
    """Write plain text content (e.g. Python scripts)."""

    safe_name = sanitize_filename(filename)
    path = directory / safe_name
    path.write_text(content, encoding="utf-8")
    return path


def build_manifest(
    request_id: str,
    product_name: str,
    directory: Path,
    files: Iterable[GeneratedFile],
) -> Manifest:
    """Create a manifest object and persist it alongside the artefacts."""

    manifest = Manifest(
        request_id=request_id,
        product_name=product_name,
        created_at=datetime.utcnow(),
        files=list(files),
    )
    manifest_path = directory / "manifest.json"
    manifest_path.write_text(
        json.dumps(manifest.model_dump(mode="json"), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return manifest


def load_manifest(request_id: str, product_name: str) -> Manifest:
    """Load a stored manifest from disk."""

    settings = get_settings()
    directory = settings.outputs_dir / _directory_name(product_name, request_id)
    if not directory.exists():
        raise FileNotFoundError("Manifest directory does not exist.")
    return load_manifest_from_directory(directory)


def load_manifest_from_directory(directory: Path) -> Manifest:
    """Load manifest using the provided directory path."""

    manifest_path = directory / "manifest.json"
    content = manifest_path.read_text(encoding="utf-8")
    manifest_data = json.loads(content)
    manifest = Manifest.model_validate(manifest_data)
    resolved_files: List[GeneratedFile] = []
    for file_meta in manifest.files:
        resolved_path = directory / Path(file_meta.path).name
        preview_url = getattr(file_meta, "preview_url", None)
        if not preview_url:
            preview_url = f"/preview/{manifest.request_id}/{file_meta.filename}"
        resolved_files.append(
            GeneratedFile(
                filename=file_meta.filename,
                path=resolved_path,
                download_url=file_meta.download_url,
                preview_url=preview_url,
                description=file_meta.description,
            )
        )
    return Manifest(
        request_id=manifest.request_id,
        product_name=manifest.product_name,
        created_at=manifest.created_at,
        files=resolved_files,
    )


def clean_expired_outputs(now: datetime | None = None) -> int:
    """Remove output directories older than the retention window."""

    settings = get_settings()
    now = now or datetime.utcnow()
    threshold = now - timedelta(hours=settings.max_output_age_hours)
    removed = 0
    for item in settings.outputs_dir.iterdir():
        if not item.is_dir():
            continue
        created_ts = datetime.utcfromtimestamp(item.stat().st_mtime)
        if created_ts < threshold:
            shutil.rmtree(item, ignore_errors=True)
            removed += 1
    return removed


def _parse_markdown_blocks(content: str) -> Iterable[tuple[str, object]]:
    table_buffer: List[List[str]] = []

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            yield ("table", table_buffer)
            table_buffer = []

    for raw_line in content.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()
        if "|" in stripped and stripped.count("|") >= 2:
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if _is_table_divider_row(cells):
                continue
            table_buffer.append(cells)
            continue

        if table_buffer:
            for item in flush_table():
                yield item

        if not stripped:
            yield ("paragraph", "")
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip() or " "
            yield ("heading", (max(1, min(level, 4)), text))
            continue

        if stripped.startswith(('- ', '* ')):
            yield ("list", stripped[2:].strip())
            continue

        match = re.match(r"(\d+)\.\s+(.*)", stripped)
        if match:
            yield ("numbered", match.group(2).strip())
            continue

        yield ("paragraph", stripped)

    if table_buffer:
        yield ("table", table_buffer)


def _is_table_divider_row(cells: Sequence[str]) -> bool:
    return all(set(cell) <= {"-", ":"} for cell in cells if cell)
